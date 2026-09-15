"""Create a marked DOCX without rebuilding rich-text paragraphs.

This deliberately fails on paragraphs containing hyperlinks, drawings, or fields.
Keep exact target keywords in ordinary text paragraphs, then highlight only those
paragraphs. It prevents the duplicate-text regression caused by rebuilding
`paragraph.text` in a document with rich OOXML content.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def has_rich_content(paragraph):
    xml = paragraph._p.xml
    return any(marker in xml for marker in ("<w:hyperlink", "<w:drawing", "<w:fldChar", "<w:instrText"))


def apply_visible_yellow_mark(run):
    """Use both Word highlight and yellow run shading for cross-viewer visibility."""
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    run.font.bold = True
    rpr = run._element.get_or_add_rPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "FFF200")
    rpr.append(shading)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("clean_docx", type=Path)
    parser.add_argument("terms_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--style", help="Only highlight paragraphs with this exact Word style name")
    args = parser.parse_args()
    terms = json.loads(args.terms_json.read_text(encoding="utf-8"))
    pattern = re.compile("(" + "|".join(re.escape(x) for x in sorted(terms, key=len, reverse=True)) + ")", re.I)
    doc = Document(args.clean_docx)
    original = "\n".join(p.text for p in paragraphs(doc))
    for paragraph in paragraphs(doc):
        if args.style and (not paragraph.style or paragraph.style.name != args.style):
            continue
        if not pattern.search(paragraph.text):
            continue
        if has_rich_content(paragraph):
            raise RuntimeError(f"Refusing to rewrite rich-content paragraph: {paragraph.text[:80]}")
        text = paragraph.text
        paragraph.clear()
        for piece in pattern.split(text):
            if not piece:
                continue
            run = paragraph.add_run(piece)
            if pattern.fullmatch(piece):
                apply_visible_yellow_mark(run)
    marked = "\n".join(p.text for p in paragraphs(doc))
    if norm(original) != norm(marked):
        raise RuntimeError("Text-integrity failure: highlighting changed document text")
    doc.save(args.output_docx)


if __name__ == "__main__":
    sys.exit(main())
