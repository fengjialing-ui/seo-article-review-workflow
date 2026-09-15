"""Create a marked review DOCX from an explicit editorial highlight map.

Use this instead of globally highlighting every occurrence of a workbook term. The map
limits yellow to the approved, reader-facing keyword optimizations and verifies that
the marked copy retains identical normalized text.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def mark_yellow(run):
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    run.font.bold = True
    rpr = run._element.get_or_add_rPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:color"), "auto")
    shade.set(qn("w:fill"), "FFF200")
    rpr.append(shade)


def main(clean_docx: str, targets_json: str, marked_docx: str):
    targets = json.loads(Path(targets_json).read_text(encoding="utf-8"))
    doc = Document(clean_docx)
    original = "\n".join(p.text for p in doc.paragraphs)
    matched = set()
    for paragraph in doc.paragraphs:
        for target in targets:
            start = target["starts_with"]
            if not norm(paragraph.text).startswith(norm(start)):
                continue
            if any(token in paragraph._p.xml for token in ("<w:drawing", "<w:hyperlink", "<w:fldChar", "<w:instrText")):
                raise RuntimeError(f"Target paragraph has rich OOXML and cannot be safely rewritten: {start}")
            terms = target["terms"]
            pattern = re.compile("(" + "|".join(re.escape(term) for term in sorted(terms, key=len, reverse=True)) + ")", re.I)
            text = paragraph.text
            if not pattern.search(text):
                raise RuntimeError(f"Approved highlight term is absent: {start}")
            paragraph.clear()
            for part in pattern.split(text):
                if not part:
                    continue
                run = paragraph.add_run(part)
                if pattern.fullmatch(part):
                    mark_yellow(run)
            matched.add(start)
            break
    missing = [target["starts_with"] for target in targets if target["starts_with"] not in matched]
    if missing:
        raise RuntimeError("Unmatched highlight targets: " + "; ".join(missing))
    marked = "\n".join(p.text for p in doc.paragraphs)
    if norm(original) != norm(marked):
        raise RuntimeError("Text-integrity failure: highlight operation changed text")
    doc.save(marked_docx)


if __name__ == "__main__":
    main(*sys.argv[1:])
