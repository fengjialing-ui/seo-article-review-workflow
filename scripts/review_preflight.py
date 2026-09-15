"""Fail-closed release validator for SEO article-review packages.

Usage:
  python scripts/review_preflight.py release.json --out preflight.json
"""
from __future__ import annotations

import argparse
from datetime import date
import hashlib
import json
import re
import sys
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from openpyxl import load_workbook
from delivery_integrity import check_render, docx_content, read_json, local_file, check_contract
from workflow_state import validate as validate_state


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


BOOTSTRAP_REQUIRED_FILES = set(read_json(Path(__file__).resolve().parents[1] / "workflow.json")["bootstrap_files"])


def workflow_bootstrap_failures(cfg, evidence):
    """Verify actual loaded-file hashes. A prefilled assertion is not evidence."""
    try:
        root = Path(cfg['workflow_root']).resolve()
        if root != Path(__file__).resolve().parents[1]:
            return ['workflow_root must identify the actual installed workflow running this validator']
        state_path = Path(cfg['workflow_state']).resolve()
        config = read_json(root/'workflow.json')
        if config['version'] != cfg.get('workflow_version'):
            return ['Workflow version differs from the canonical workflow.json']
        return validate_state(read_json(state_path), config, root, state_path.parent, complete=False)
    except (KeyError, ValueError, TypeError, OSError) as exc:
        return ['Missing/invalid generated workflow bootstrap: '+str(exc)]


def doc_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def canonical_terms(path: Path, expected_sheet: str | None):
    workbook = load_workbook(path, read_only=True, data_only=True)
    choices = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        for row_no, row in enumerate(rows):
            cells = [norm(str(value)) if value is not None else "" for value in row]
            if "keyword" not in cells or "search volume" not in cells:
                continue
            key_col = cells.index("keyword")
            volume_col = cells.index("search volume")
            terms = []
            for data in rows[row_no + 1:]:
                term = data[key_col] if key_col < len(data) else None
                volume = data[volume_col] if volume_col < len(data) else None
                if isinstance(term, str) and term.strip() and isinstance(volume, (int, float)):
                    terms.append(term.strip())
            choices.append((sheet.title, terms))
    if not choices:
        raise ValueError("No canonical Keyword + Search Volume sheet found")
    if expected_sheet:
        matches = [terms for title, terms in choices if title == expected_sheet]
        if not matches:
            raise ValueError(f"Configured keyword sheet not found: {expected_sheet}")
        return expected_sheet, matches[0]
    title, terms = max(choices, key=lambda item: len(item[1]))
    return title, terms


def media_count(docx: Path) -> int:
    with zipfile.ZipFile(docx) as archive:
        return len([name for name in archive.namelist() if name.startswith("word/media/")])


def media_hashes(docx: Path) -> set[str]:
    with zipfile.ZipFile(docx) as archive:
        return {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def word_count(doc) -> int:
    return len(re.findall(r"\b[\w'-]+\b", " ".join(p.text for p in doc_paragraphs(doc))))


def document_hygiene_failures(docx: Path, clean_copy: bool) -> list[str]:
    failures = []
    with zipfile.ZipFile(docx) as archive:
        names = set(archive.namelist())
        if clean_copy and "word/comments.xml" in names:
            failures.append(f"Clean copy contains Word comments: {docx.name}")
        xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
        if clean_copy and ("<w:ins" in xml or "<w:del" in xml):
            failures.append(f"Clean copy contains tracked changes: {docx.name}")
        if re.search(r"\[(?:IMAGE|TODO|TBD|INSERT|PLACEHOLDER)\]", xml, flags=re.I):
            failures.append(f"Placeholder marker found in document XML: {docx.name}")
    return failures


def section_has_image(doc, heading: str) -> bool:
    paragraphs = list(doc.paragraphs)
    start = next((i for i, p in enumerate(paragraphs) if p.text.strip() == heading), None)
    if start is None:
        return False
    for paragraph in paragraphs[start + 1:]:
        if paragraph.style and paragraph.style.name in {"Heading 1", "Heading 2"}:
            break
        if "<w:drawing" in paragraph._p.xml:
            return True
    return False


def nonempty(value) -> bool:
    return bool(str(value or "").strip())


def inside(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def v35_evidence_failures(evidence, h2_headings, heading_rows_actual, source_words, final_words):
    """Validate the non-bypassable v3.5 editorial controls."""
    failures = []
    economy = evidence.get("content_economy") or {}
    for field in ("source_word_count", "final_word_count", "duplicate_decision_rows"):
        if economy.get(field) in (None, "", []):
            failures.append(f"Content-economy evidence missing: {field}")
    if economy:
        try:
            if int(economy.get("source_word_count")) != source_words or int(economy.get("final_word_count")) != final_words:
                failures.append("Content-economy word counts do not match the actual source/final DOCX")
        except (TypeError, ValueError):
            failures.append("Content-economy word counts are invalid")
        if final_words > source_words * 1.10 and not nonempty(economy.get("expansion_justification")):
            failures.append("Final article exceeds 110% of source words without a section-by-section expansion justification")
        for row in economy.get("duplicate_decision_rows", []):
            missing = [key for key in ("decision", "primary_owner", "duplicate_disposition", "reader_value_after_edit") if not nonempty(row.get(key))]
            if missing:
                failures.append("Incomplete duplicate-decision row: " + ", ".join(missing))

    heading_rows = evidence.get("heading_rows") or []
    mapped = {(norm(row.get("heading")), norm(row.get("level"))): row for row in heading_rows if nonempty(row.get("heading"))}
    for heading, level in heading_rows_actual:
        row = mapped.get((norm(heading), norm(level)))
        if not row:
            failures.append(f"Heading-role map lacks {level}: {heading}")
            continue
        role = norm(row.get("role"))
        if role not in {"primary_keyword", "semantic_keyword", "long_tail_question", "action_step", "reference", "structural"}:
            failures.append(f"Heading-role map has invalid role for {heading}: {row.get('role')}")
        if not nonempty(row.get("naturalness")):
            failures.append(f"Heading-role map has no naturalness decision for {heading}")
        if level in {"H1", "H2"} and role not in {"reference", "structural"} and not nonempty(row.get("keyword_or_semantic_cluster")):
            failures.append(f"Core {level} has no keyword/semantic cluster: {heading}")
        if level == "H3" and role == "action_step" and not nonempty(row.get("action_parent_rationale")):
            failures.append(f"Action-led H3 has no parent-intent rationale: {heading}")

    source_rows = evidence.get("source_records") or []
    if not source_rows:
        failures.append("No deduplicated claim-linked source records")
    urls = set()
    for row in source_rows:
        missing = [key for key in ("claim_scope", "official_url", "verified_date", "final_location") if not nonempty(row.get(key))]
        if missing:
            failures.append("Incomplete source record: " + ", ".join(missing))
            continue
        url = norm(row["official_url"])
        if not url.startswith("https://"):
            failures.append("Source record URL is not HTTPS: " + row["official_url"])
        if url in urls:
            failures.append("Duplicate source URL in claim-linked source records: " + row["official_url"])
        urls.add(url)
        try:
            date.fromisoformat(str(row["verified_date"]))
        except ValueError:
            failures.append("Invalid source verification date: " + str(row["verified_date"]))

    for product in evidence.get("product_rows", []):
        for field in ("availability_status", "availability_verified_date", "availability_action"):
            if not nonempty(product.get(field)):
                failures.append(f"Product availability evidence missing for {product.get('name', '<unnamed>')}: {field}")
        try:
            date.fromisoformat(str(product.get("availability_verified_date", "")))
        except ValueError:
            failures.append(f"Invalid product availability verification date for {product.get('name', '<unnamed>')}")
    return failures


def core_conversion_ultra_tips_failures(cfg, products, clean_text, visual_uses):
    """Fail closed when a named conversion product has an invalid route.

    The delivery contract supplies the classification, but the article and release
    evidence must prove that the declared route is internally consistent. A direct
    product may never retain an Ultra Tips/adjacent route from an older template.
    """
    failures = []
    requested_cores = cfg.get("core_conversion_products", [])
    if not isinstance(requested_cores, list):
        return ["core_conversion_products must be a list in the release manifest"]

    requested_names = {norm(name) for name in requested_cores if nonempty(name)}
    product_names = {norm(row.get("name")) for row in products if nonempty(row.get("name"))}
    for name in requested_names - product_names:
        failures.append(f"Named core conversion product is missing from product_rows: {name}")

    visual_assets = {str(use.get("asset_filename", "")).strip() for use in visual_uses}
    valid_support = {"direct_supported", "adjacent_authorized_route", "no_truthful_route"}
    required_touchpoint_roles = {"decision_context", "ultra_tips_owner", "reassurance"}
    blocked_cta_pattern = re.compile(r"\b(?:try|start|use|upload|download)\s+(?:now|today)\b", re.I)

    for row in products:
        name = row.get("name", "<unnamed>")
        is_core = bool(row.get("core_conversion_product")) or norm(name) in requested_names
        if not is_core:
            continue
        if not row.get("core_conversion_product"):
            failures.append(f"Named core conversion product is not marked core_conversion_product: {name}")

        support = norm(row.get("primary_task_support"))
        if support not in valid_support:
            failures.append(f"Core conversion product has invalid primary_task_support for {name}")
            continue
        if support == "direct_supported":
            if norm(row.get("role")) not in {"direct solution", "direct", "required"}:
                failures.append(f"Direct-supported core product is not classified as a direct solution: {name}")
            direct_heading = row.get("direct_solution_heading")
            if not nonempty(direct_heading):
                failures.append(f"Direct-supported core product lacks direct_solution_heading evidence: {name}")
            elif norm(direct_heading) not in norm(clean_text) or norm(name) not in norm(direct_heading):
                failures.append(f"Direct-supported core product heading is not evidenced in the clean article: {name}")
            direct_steps = row.get("direct_workflow_steps")
            if not isinstance(direct_steps, list) or not 3 <= len(direct_steps) <= 6 or any(not nonempty(step) for step in direct_steps):
                failures.append(f"Direct-supported core product requires three to six direct workflow steps: {name}")
            else:
                for step in direct_steps:
                    if norm(step) not in norm(clean_text):
                        failures.append(f"Direct-supported core product workflow step is not evidenced in the clean article: {name}")
                        break
            official_url = str(row.get("official_evidence_url", "")).strip()
            if not official_url or norm(official_url) not in norm(clean_text):
                failures.append(f"Direct-supported core product lacks a traceable official evidence URL in the clean article: {name}")
            module = row.get("ultra_tips")
            if module not in (None, {}, []):
                failures.append(f"Direct-supported core product retains an Ultra Tips evidence object: {name}")
            escaped_name = re.escape(str(name))
            if re.search(rf"ultra\s+tips[^\n.]{{0,160}}{escaped_name}|{escaped_name}[^\n.]{{0,160}}ultra\s+tips", clean_text, re.I):
                failures.append(f"Direct-supported core product is routed through Ultra Tips in the clean article: {name}")
            continue
        if support == "no_truthful_route":
            if not nonempty(row.get("blocked_decision")):
                failures.append(f"Core conversion product has no truthful route but no blocked_decision evidence: {name}")
            if not nonempty(row.get("approved_alternative_request")):
                failures.append(f"Core conversion product has no truthful route but no approved_alternative_request: {name}")
            failures.append(f"Core conversion product route is blocked pending approved alternative: {name}")
            continue

        # Any other route is a required, detailed Boundary-sensitive core conversion module.
        if norm(row.get("role")) not in {"boundary-sensitive core conversion", "boundary-sensitive core conversion (ultra tips)"}:
            failures.append(f"Core conversion product with adjacent route is not classified as Boundary-sensitive core conversion (Ultra Tips): {name}")
        module = row.get("ultra_tips") or {}
        fields = [
            "heading", "boundary_disclosure", "supported_authorized_route", "not_for",
            "current_status_statement", "cta_text", "visual_asset_filename",
        ]
        missing = [field for field in fields if not nonempty(module.get(field))]
        if missing:
            failures.append(f"Ultra Tips evidence incomplete for {name}: missing {', '.join(missing)}")
            continue
        if "ultra tips" not in norm(module["heading"]) or norm(name) not in norm(module["heading"]):
            failures.append(f"Ultra Tips heading must name both Ultra Tips and the core product: {name}")
        steps = module.get("workflow_steps")
        if not isinstance(steps, list) or not 4 <= len(steps) <= 6 or any(not nonempty(step) for step in steps):
            failures.append(f"Ultra Tips workflow must have four to six non-empty ordered steps: {name}")
        touchpoints = module.get("touchpoints")
        if not isinstance(touchpoints, list) or len(touchpoints) < 3:
            failures.append(f"Ultra Tips requires three touchpoints for {name}")
        else:
            roles = {norm(point.get("role")) for point in touchpoints if isinstance(point, dict)}
            if not required_touchpoint_roles.issubset(roles):
                failures.append(f"Ultra Tips touchpoint roles missing for {name}: decision context, owner, and reassurance are required")
            for point in touchpoints:
                if not isinstance(point, dict) or not nonempty(point.get("exact_location")):
                    failures.append(f"Ultra Tips touchpoint lacks an exact location for {name}")
                    break
                if norm(point["exact_location"]) not in norm(clean_text):
                    failures.append(f"Ultra Tips touchpoint location is not evidenced in the clean article for {name}: {point['exact_location']}")
                    break
        for text_field in ("heading", "boundary_disclosure", "supported_authorized_route", "not_for", "current_status_statement", "cta_text"):
            text = str(module[text_field])
            if norm(text) not in norm(clean_text):
                failures.append(f"Ultra Tips {text_field} is not evidenced in the clean article for {name}")
        if str(module["visual_asset_filename"]).strip() not in visual_assets:
            failures.append(f"Ultra Tips visual asset is not mapped in visual_uses for {name}")
        elif not any(
            str(use.get("asset_filename", "")).strip() == str(module["visual_asset_filename"]).strip()
            and norm(use.get("required_section")) == norm(module["heading"])
            for use in visual_uses
        ):
            failures.append(f"Ultra Tips visual asset is not mapped to the Ultra Tips H2 for {name}")
        degraded = norm(row.get("availability_status")) in {"coming_soon", "beta", "demo_only", "unavailable", "region_limited"}
        if degraded and blocked_cta_pattern.search(str(module["cta_text"])):
            failures.append(f"Degraded product availability has a conventional immediate-use CTA in Ultra Tips: {name}")
    return failures


def evidence_failures(cfg, top20, clean_text, yellow, asset_dir: Path, required_sections):
    """Validate specific release evidence, not only the presence of record files."""
    evidence_path = cfg.get("release_evidence_json")
    if not evidence_path:
        return ["Missing release_evidence_json: v3.4 requires specific, machine-readable evidence"], {}
    path = Path(evidence_path)
    if not path.is_file():
        return [f"Release evidence file not found: {path}"], {}
    try:
        evidence = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        return [f"Invalid release evidence JSON: {exc}"], {}

    failures = []
    if evidence.get("release_id") != cfg.get("release_id"):
        failures.append("Release evidence ID does not match manifest release_id")
    failures.extend(workflow_bootstrap_failures(cfg, evidence))
    if not evidence.get("source_audit_complete"):
        failures.append("Source asset audit is not complete in release evidence")
    if not evidence.get("structure_plan_complete"):
        failures.append("SERP-backed structure plan is not complete")
    if evidence.get("reference_outline_copied"):
        failures.append("Reference outline was copied; sample articles may not prescribe structure")
    try:
        date.fromisoformat(str(evidence.get("serp_research_date", "")))
    except ValueError:
        failures.append("SERP research date is missing or invalid")
    try:
        retention = float(evidence.get("source_retention_pct"))
        if retention < 80 and not nonempty(evidence.get("retention_authorization")):
            failures.append(f"Effective-information retention {retention}% is below 80% without explicit authorization")
    except (TypeError, ValueError):
        failures.append("Source retention percentage is missing or invalid")
    try:
        source_assets = int(evidence.get("source_assets_total"))
        dispositioned = int(evidence.get("source_assets_dispositioned"))
        protected = int(evidence.get("protected_assets_total"))
        protected_kept = int(evidence.get("protected_assets_preserved_or_equivalent"))
        if source_assets < 1 or dispositioned != source_assets:
            failures.append("Source audit is not asset-complete: every source asset needs a disposition")
        if protected < 1 or protected_kept != protected:
            failures.append("Protected source assets were not all preserved or replaced with documented equivalents")
    except (TypeError, ValueError):
        failures.append("Source-audit counts are missing or invalid")

    rows = {norm(row.get("keyword")): row for row in evidence.get("keyword_rows", []) if nonempty(row.get("keyword"))}
    used_exception_reasons = {}
    for term in top20:
        row = rows.get(norm(term))
        if not row:
            failures.append(f"Top-20 keyword lacks a release-evidence row: {term}")
            continue
        exception = row.get("exception")
        if exception:
            fields = ["reason", "semantic_route", "reader_facing_substitute", "approved_by"]
            missing = [field for field in fields if not nonempty(exception.get(field))]
            if missing:
                failures.append(f"Incomplete exception for {term}: missing {', '.join(missing)}")
            reason = norm(exception.get("reason"))
            if reason:
                used_exception_reasons.setdefault(reason, []).append(term)
            continue
        fields = ["intent", "optimization_role", "exact_location", "reader_value", "highlighted_text", "context_text"]
        missing = [field for field in fields if not nonempty(row.get(field))]
        if missing:
            failures.append(f"Incomplete keyword evidence for {term}: missing {', '.join(missing)}")
            continue
        highlighted = norm(row["highlighted_text"])
        if highlighted not in norm(clean_text) or highlighted not in norm(yellow):
            failures.append(f"Keyword highlight evidence does not match final clean/marked text: {term}")
        if any(token in norm(row["exact_location"]) for token in ("contextual article sentence", "keyword coverage", "tbd")):
            failures.append(f"Keyword location is generic rather than exact: {term}")
        if norm(row["context_text"]) == norm(row["highlighted_text"]) or len(str(row["context_text"]).strip()) < len(str(row["highlighted_text"]).strip()) + 15:
            failures.append(f"Keyword evidence has no useful reader-facing context: {term}")
    for reason, terms in used_exception_reasons.items():
        if len(terms) > 1:
            failures.append("Generic keyword exception reason reused for: " + "; ".join(terms))
    required_roles = {norm(role) for role in cfg.get("required_keyword_roles", [])}
    actual_roles = {norm(row.get("optimization_role")) for row in rows.values() if not row.get("exception")}
    missing_roles = required_roles - actual_roles
    if missing_roles:
        failures.append("Required keyword optimization roles are missing: " + "; ".join(sorted(missing_roles)))

    products = evidence.get("product_rows", [])
    if cfg.get("product_required", True) and not products:
        failures.append("Product evidence is required but product_rows is empty")
    for row in products:
        fields = ["name", "role", "differentiated_value", "exact_location", "official_evidence_url", "cta", "constraint"]
        missing = [field for field in fields if not nonempty(row.get(field))]
        if len(row.get("best_for_scenarios", [])) < 3 or any(not nonempty(x) for x in row.get("best_for_scenarios", [])):
            missing.append("three best-for scenarios")
        if not row.get("not_best_for") or any(not nonempty(x) for x in row.get("not_best_for", [])):
            missing.append("not-best-for / escalation")
        if missing:
            failures.append(f"Incomplete product evidence for {row.get('name', '<unnamed>')}: missing {', '.join(missing)}")
        elif not str(row["official_evidence_url"]).startswith("https://"):
            failures.append(f"Product evidence URL is not a current HTTPS official URL: {row['name']}")

    failures.extend(core_conversion_ultra_tips_failures(cfg, products, clean_text, evidence.get("visual_uses", [])))

    uses_by_section = {}
    visual_fields = ["asset_filename", "reader_persona_decision", "visual_job", "concrete_scene", "style_rationale", "rights_status", "alt_text", "caption", "exact_insertion_point", "rendered_qa"]
    for use in evidence.get("visual_uses", []):
        section = use.get("required_section")
        if nonempty(section):
            uses_by_section.setdefault(section, []).append(use)
        missing = [field for field in visual_fields if not nonempty(use.get(field))]
        if missing:
            failures.append(f"Incomplete visual-use evidence for section {section or '<unnamed>'}: missing {', '.join(missing)}")
        elif not (asset_dir / use["asset_filename"]).is_file():
            failures.append(f"Visual-use asset missing from asset directory: {use['asset_filename']}")
        elif norm(use["rendered_qa"]) != "pass":
            pending_ok = (
                cfg.get("render_status") == "STRUCTURALLY_VERIFIED_VISUAL_QA_PENDING"
                and norm(use["rendered_qa"]) == "structurally verified - visual qa pending"
            )
            if not pending_ok:
                failures.append(f"Visual-use lacks rendered QA PASS: {use['asset_filename']}")
    for section in required_sections:
        count = len(uses_by_section.get(section, []))
        if count < 1 or count > 2:
            failures.append(f"Section '{section}' has {count} visual-use rows; contract requires one or two")

    human_qa = evidence.get("editorial_qa") or {}
    qa_fields = [
        "content_integrity", "keyword_editorial_quality", "product_suitability",
        "visual_relevance", "visual_aesthetics", "sources_checked",
        "document_hygiene", "release_purity",
    ]
    if not nonempty(human_qa.get("reviewer_name")):
        failures.append("Final editorial QA has no named reviewer")
    if human_qa.get('reviewer_type') not in {'ai', 'human'}:
        failures.append('Editorial QA must accurately declare reviewer_type: ai or human')
    if cfg.get('human_approval_required') is True:
        approval = evidence.get('human_approval') or {}
        if approval.get('reviewer_type') != 'human' or approval.get('result') != 'pass' or not approval.get('reviewer_name') or not approval.get('approval_reference'):
            failures.append('Required real human approval is missing')
    elif cfg.get('human_approval_required') is not False:
        failures.append('Declare human_approval_required explicitly from the task/project contract')
    try:
        date.fromisoformat(str(human_qa.get("review_date", "")))
    except ValueError:
        failures.append("Final editorial QA date is missing or invalid")
    blocked_qa = [field for field in qa_fields if norm(human_qa.get(field)) != "pass"]
    if blocked_qa:
        failures.append("Final editorial QA has unresolved criteria: " + ", ".join(blocked_qa))
    if cfg.get("render_status") == "PASS" and not human_qa.get("rendered_pages_reviewed"):
        failures.append("Manifest claims render PASS but final editorial QA does not confirm rendered pages were reviewed")

    if cfg.get("eeat_required", True) and not [source for source in evidence.get("eeat_sources_checked", []) if nonempty(source)]:
        failures.append("No EEAT/official sources were recorded")
    if not evidence.get("release_directory_pure"):
        failures.append("Release evidence does not confirm single-release directory purity")
    release_dir = cfg.get("release_dir")
    if not release_dir:
        failures.append("Manifest does not declare release_dir for purity validation")
    else:
        root = Path(release_dir)
        if not root.is_dir():
            failures.append(f"Configured release_dir does not exist: {root}")
            return failures, evidence
        for path_value in [cfg["clean_docx"], cfg["marked_docx"], *cfg.get("supporting_files", [])]:
            if not inside(Path(path_value), root):
                failures.append(f"Release artifact sits outside configured release_dir: {path_value}")
    return failures, evidence


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    cfg = json.loads(args.manifest.read_text(encoding="utf-8"))
    # Resolve config paths against the manifest rather than the caller's cwd.
    manifest_root = args.manifest.resolve().parent
    for field in ['source_docx','clean_docx','marked_docx','keyword_xlsx','asset_dir','release_dir','release_evidence_json','workflow_root','workflow_state']:
        if cfg.get(field): cfg[field] = str((manifest_root / cfg[field]).resolve())
    for field in ['required_record_files','supporting_files']:
        if isinstance(cfg.get(field), list): cfg[field] = [str((manifest_root / x).resolve()) for x in cfg[field]]
    required_manifest_fields = ["release_id", "workflow_version", "source_docx", "clean_docx", "marked_docx", "keyword_xlsx", "asset_dir", "release_dir", "release_evidence_json", "expected_release_files", "required_keyword_roles", "required_record_files", "visual_scope"]
    missing_manifest_fields = [field for field in required_manifest_fields if not nonempty(cfg.get(field))]
    if missing_manifest_fields:
        result = {
            "release_id": cfg.get("release_id"),
            "checks": {"manifest_schema": False},
            "errors": ["Manifest is missing required fields: " + ", ".join(missing_manifest_fields)],
            "status": "FAIL",
            "publication_ready": False,
        }
        encoded = json.dumps(result, ensure_ascii=False, indent=2)
        if args.out:
            args.out.write_text(encoded, encoding="utf-8")
        print(encoded)
        return 1
    malformed_lists = [
        field for field in ("expected_release_files", "required_keyword_roles", "required_record_files")
        if not isinstance(cfg.get(field), list) or not cfg[field]
    ]
    if malformed_lists:
        result = {
            "release_id": cfg.get("release_id"),
            "checks": {"manifest_schema": False},
            "errors": ["Manifest requires non-empty lists: " + ", ".join(malformed_lists)],
            "status": "FAIL",
            "publication_ready": False,
        }
        encoded = json.dumps(result, ensure_ascii=False, indent=2)
        if args.out:
            args.out.write_text(encoded, encoding="utf-8")
        print(encoded)
        return 1
    clean_path = Path(cfg["clean_docx"])
    marked_path = Path(cfg["marked_docx"])
    source_path = Path(cfg["source_docx"])
    if not source_path.is_file():
        raise FileNotFoundError(f"Source DOCX not found: {source_path}")
    if not clean_path.is_file() or not marked_path.is_file():
        raise FileNotFoundError("Clean or marked DOCX not found")
    source = Document(source_path)
    clean = Document(clean_path)
    marked = Document(marked_path)
    clean_data = docx_content(clean_path)
    marked_data = docx_content(marked_path)
    clean_text = clean_data['text']
    marked_text = marked_data['text']
    sheet, terms = canonical_terms(Path(cfg["keyword_xlsx"]), cfg.get("keyword_sheet"))
    top20 = terms[:20]
    exceptions = {norm(item["keyword"]): item for item in cfg.get("top20_exceptions", [])}
    errors, checks = [], {}
    try:
        state_path=Path(cfg['workflow_state'])
        config=read_json(Path(cfg['workflow_root'])/'workflow.json')
        state=read_json(state_path)
        errors.extend(validate_state(state,config,Path(cfg['workflow_root']),state_path.parent))
        errors.extend(check_contract(cfg,state_path.parent,config['contract_fields']))
        if state.get('run_id') != cfg.get('release_id'): errors.append('Workflow state belongs to another release')
    except (KeyError,ValueError,OSError,TypeError) as exc:
        errors.append('Workflow stage evidence invalid: '+str(exc))

    h1_count = sum(1 for p in doc_paragraphs(clean) if p.style and p.style.name == "Heading 1")
    checks["one_h1"] = h1_count == 1
    if h1_count != 1:
        errors.append(f"Expected one H1, got {h1_count}")

    source_words, final_words = word_count(source), word_count(clean)
    word_ratio = final_words / max(source_words, 1)
    checks["source_word_count_regression"] = word_ratio >= 0.8 or nonempty(cfg.get("shortening_authorization"))
    if not checks["source_word_count_regression"]:
        errors.append(f"Final article has {final_words}/{source_words} words ({word_ratio:.1%}); below 80% without authorization")

    checks["marked_clean_text_identical"] = norm(clean_text) == norm(marked_text)
    if not checks["marked_clean_text_identical"]:
        errors.append("Marked and clean text differ after normalization")
    checks["marked_clean_media_identical"] = media_count(clean_path) == media_count(marked_path) and clean_data['media'] == marked_data['media']
    if not checks["marked_clean_media_identical"]:
        errors.append("Marked and clean media counts differ")
    if clean_data['links'] != marked_data['links']:
        errors.append('Marked and clean hyperlink targets differ')
    errors.extend(document_hygiene_failures(clean_path, clean_copy=True))
    errors.extend(document_hygiene_failures(marked_path, clean_copy=False))

    yellow = "\n".join(
        run.text for paragraph in doc_paragraphs(marked) for run in paragraph.runs
        if run.font.highlight_color == WD_COLOR_INDEX.YELLOW
    )
    keyword_failures = []
    for term in top20:
        if norm(term) in exceptions:
            item = exceptions[norm(term)]
            if not str(item.get("reason", "")).strip() or not str(item.get("semantic_coverage", "")).strip():
                keyword_failures.append(f"incomplete exception record: {term}")
            continue
        if norm(term) not in norm(clean_text):
            keyword_failures.append(f"missing exact text: {term}")
        elif norm(term) not in norm(yellow):
            keyword_failures.append(f"missing yellow highlight: {term}")
    checks["top20_keyword_disposition"] = not keyword_failures
    errors.extend(keyword_failures)

    images = list(clean.inline_shapes)
    missing_alt = []
    for idx, shape in enumerate(images, 1):
        if not shape._inline.docPr.get("descr", "").strip():
            missing_alt.append(idx)
    captions = [p for p in doc_paragraphs(clean) if p.style and "caption" in p.style.name.casefold()]
    checks["image_alt"] = not missing_alt
    checks["image_captions"] = len(captions) >= len(images)
    if missing_alt:
        errors.append(f"Images missing alt text: {missing_alt}")
    if len(captions) < len(images):
        errors.append(f"Caption count {len(captions)} < image count {len(images)}")

    h2_headings = [p.text.strip() for p in clean.paragraphs if p.style and p.style.name == "Heading 2"]
    heading_rows_actual = [
        (p.text.strip(), {"Heading 1": "H1", "Heading 2": "H2", "Heading 3": "H3"}[p.style.name])
        for p in clean.paragraphs
        if p.style and p.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]
    if cfg.get("visual_scope") == "ALL_H2":
        exceptions = cfg.get("visual_section_exceptions", [])
        exception_names = set()
        for item in exceptions:
            if not isinstance(item, dict) or not nonempty(item.get("heading")) or not nonempty(item.get("reason")) or not nonempty(item.get("approved_by")):
                errors.append("Every ALL_H2 visual exception needs heading, reason, and approved_by")
                continue
            exception_names.add(item["heading"].strip())
        unknown = exception_names - set(h2_headings)
        if unknown:
            errors.append("Visual exception names a non-existent H2: " + "; ".join(sorted(unknown)))
        required_sections = [heading for heading in h2_headings if heading not in exception_names]
    elif cfg.get("visual_scope") == "NAMED_SECTIONS":
        required_sections = cfg.get("required_visual_sections", [])
        if not required_sections and cfg.get('visuals_not_required') is not True:
            errors.append('Named visual scope needs sections or an explicit no-visual task contract')
        if set(required_sections) - set(h2_headings):
            errors.append('Named visual scope contains a non-existent H2')
    else:
        errors.append("visual_scope must be ALL_H2 or NAMED_SECTIONS")
        required_sections = cfg.get("required_visual_sections", [])
    missing_sections = [heading for heading in required_sections if not section_has_image(clean, heading)]
    checks["per_section_visual_coverage"] = not missing_sections
    if missing_sections:
        errors.append("Required sections without an image: " + "; ".join(missing_sections))

    asset_dir = Path(cfg["asset_dir"])
    expected_assets = cfg.get("asset_files", [])
    absent_assets = [name for name in expected_assets if not (asset_dir / name).is_file()]
    checks["independent_assets"] = not absent_assets
    if absent_assets:
        errors.append("Missing independent assets: " + "; ".join(absent_assets))

    evidence_errors, evidence = evidence_failures(
        cfg, top20, clean_text, yellow, asset_dir, required_sections
    )
    evidence_errors.extend(v35_evidence_failures(evidence, h2_headings, heading_rows_actual, source_words, final_words))
    checks["release_evidence"] = not evidence_errors
    errors.extend(evidence_errors)

    clean_media = media_hashes(clean_path)
    visual_hash_failures = []
    for use in evidence.get("visual_uses", []):
        filename = use.get("asset_filename")
        if filename and (asset_dir / filename).is_file() and file_hash(asset_dir / filename) not in clean_media:
            visual_hash_failures.append(filename)
    checks["mapped_visuals_embedded"] = not visual_hash_failures
    if visual_hash_failures:
        errors.append("Mapped independent visual files are not embedded in clean DOCX: " + "; ".join(sorted(set(visual_hash_failures))))

    release_root = Path(cfg["release_dir"])
    actual_release_files = {
        path.relative_to(release_root).as_posix()
        for path in release_root.rglob("*") if path.is_file()
    }
    expected_release_files = {str(item).replace("\\", "/") for item in cfg["expected_release_files"]}
    unexpected = actual_release_files - expected_release_files
    missing_declared = expected_release_files - actual_release_files
    checks["release_directory_exact_manifest"] = not unexpected and not missing_declared
    if unexpected:
        errors.append("Unexpected files in release directory: " + "; ".join(sorted(unexpected)))
    if missing_declared:
        errors.append("Expected release files missing: " + "; ".join(sorted(missing_declared)))

    supporting_files = [Path(path) for path in cfg.get("supporting_files", [])]
    missing_supporting = [str(path) for path in supporting_files if not path.is_file()]
    checks["supporting_records"] = not missing_supporting
    if missing_supporting:
        errors.append("Missing required supporting records: " + "; ".join(missing_supporting))
    required_records = [Path(path) for path in cfg["required_record_files"]]
    missing_records = [str(path) for path in required_records if not path.is_file()]
    checks["named_required_records"] = not missing_records
    if missing_records:
        errors.append("Missing named required records: " + "; ".join(missing_records))

    render_status = cfg.get("render_status")
    if render_status == 'PASS':
        for name, docx in [('clean',clean_path),('marked',marked_path)]:
            try:
                report_path=local_file(Path(cfg['release_dir']),cfg.get('render_reports',{}).get(name))
                errors.extend(name+': '+e for e in check_render(read_json(report_path),Path(cfg['release_dir']),docx))
            except (ValueError,OSError,TypeError) as exc:
                errors.append(name+' render evidence: '+str(exc))
    checks["render_status_recorded"] = render_status in {"PASS", "STRUCTURALLY_VERIFIED_VISUAL_QA_PENDING"}
    if not checks["render_status_recorded"]:
        errors.append("Render status must be PASS or STRUCTURALLY_VERIFIED_VISUAL_QA_PENDING")

    result = {
        "release_id": cfg.get("release_id"),
        "canonical_keyword_sheet": sheet,
        "top20": top20,
        "source_word_count": source_words,
        "final_word_count": final_words,
        "final_to_source_word_ratio": round(word_ratio, 4),
        "checks": checks,
        "errors": errors,
        "status": "FAIL" if errors else ("PASS" if render_status == "PASS" else "STRUCTURALLY_VERIFIED_VISUAL_QA_PENDING"),
        "publication_ready": not errors and render_status == "PASS",
    }
    encoded = json.dumps(result, ensure_ascii=False, indent=2)
    if args.out:
        args.out.write_text(encoded, encoding="utf-8")
    print(encoded)
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
