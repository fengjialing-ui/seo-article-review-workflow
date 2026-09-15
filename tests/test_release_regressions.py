import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from test_execution import ROOT, SCRIPTS, png, save, complete_stages
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from openpyxl import Workbook
from delivery_integrity import digest
from review_preflight import workflow_bootstrap_failures, word_count

def package_fixture(run):
    article=Document();article.add_heading('Topic',0)
    article.paragraphs[0].style='Heading 1'
    article.add_paragraph('This useful Topic explanation preserves the original reader decision and gives a clear next action.')
    article.add_heading('Sources',level=2);article.add_paragraph('https://example.com/official')
    article.save(run/'source.docx');article.save(run/'clean.docx')
    article.paragraphs[0].runs[0].font.highlight_color=WD_COLOR_INDEX.YELLOW;article.save(run/'marked.docx')
    workbook=Workbook();sheet=workbook.active;sheet.title='Keywords';sheet.append(['Keyword','Search Volume']);sheet.append(['Topic',100]);workbook.save(run/'keywords.xlsx')
    (run/'assets').mkdir();(run/'record.md').write_text('Synthetic source, keyword and editorial evidence for a validator test.')
    png(run/'page.png')
    for name in ['clean','marked']:
        save(run/(name+'-render.json'),{'status':'pass','source_docx_sha256':digest(run/(name+'.docx')),'page_count':1,'page_files':[{'path':'page.png','sha256':digest(run/'page.png')}],'reviewer_type':'ai','reviewed_by':'Regression fixture','inspection':'Synthetic rendered-page evidence'})
    words=word_count(Document(run/'clean.docx'))
    evidence={'release_id':'TEST','source_retention_pct':100,'source_audit_complete':True,'source_assets_total':1,'source_assets_dispositioned':1,'protected_assets_total':1,'protected_assets_preserved_or_equivalent':1,'serp_research_date':'2026-09-07','structure_plan_complete':True,'reference_outline_copied':False,
      'content_economy':{'source_word_count':words,'final_word_count':words,'duplicate_decision_rows':[{'decision':'Explain topic','primary_owner':'Opening','duplicate_disposition':'No duplicates','reader_value_after_edit':'Preserves original decision'}]},
      'heading_rows':[{'heading':'Topic','level':'H1','role':'primary_keyword','keyword_or_semantic_cluster':'Topic','naturalness':'PASS'},{'heading':'Sources','level':'H2','role':'reference','naturalness':'PASS'}],
      'source_records':[{'claim_scope':'Synthetic claim','official_url':'https://example.com/official','verified_date':'2026-09-07','final_location':'Sources'}],
      'keyword_rows':[{'keyword':'Topic','intent':'Understand topic','optimization_role':'h1','exact_location':'Article H1','reader_value':'Identifies the subject','highlighted_text':'Topic','context_text':'This useful Topic explanation preserves the original reader decision.'}],
      'product_rows':[],'visual_uses':[],'eeat_sources_checked':['https://example.com/official'],'release_directory_pure':True,
      'editorial_qa':{'reviewer_type':'ai','reviewer_name':'Regression fixture','review_date':'2026-09-07','rendered_pages_reviewed':True,**{x:'PASS' for x in ['content_integrity','keyword_editorial_quality','product_suitability','visual_relevance','visual_aesthetics','sources_checked','document_hygiene','release_purity']}}}
    save(run/'evidence.json',evidence)
    complete_stages(run,'record.md')
    cfg={'release_id':'TEST','workflow_version':'3.8.0','workflow_root':str(ROOT),'workflow_state':'workflow-state.json','source_docx':'source.docx','clean_docx':'clean.docx','marked_docx':'marked.docx','keyword_xlsx':'keywords.xlsx','keyword_sheet':'Keywords','asset_dir':'assets','release_dir':'.','release_evidence_json':'evidence.json','required_keyword_roles':['h1'],'required_record_files':['record.md'],'supporting_files':[],'visual_scope':'NAMED_SECTIONS','required_visual_sections':[],'visuals_not_required':True,'product_required':False,'core_conversion_products':[],'eeat_required':True,'human_approval_required':False,'render_status':'PASS','render_reports':{'clean':'clean-render.json','marked':'marked-render.json'}}
    cfg['expected_release_files']=sorted([p.relative_to(run).as_posix() for p in run.rglob('*') if p.is_file()]+['manifest.json'])
    save(run/'manifest.json',cfg)

def run_gate(run):
    result=subprocess.run([sys.executable,str(SCRIPTS/'review_preflight.py'),str(run/'manifest.json')],capture_output=True,text=True)
    return result.returncode, result.stdout+result.stderr

class ReviewReleaseTests(unittest.TestCase):
    def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.run=Path(self.tmp.name);package_fixture(self.run)
    def tearDown(self):self.tmp.cleanup()
    def test_complete_review_passes(self):
        code,out=run_gate(self.run);self.assertEqual(code,0,out);self.assertTrue(json.loads(out)['publication_ready'])
    def test_legacy_prefilled_bootstrap_rejected(self):
        evidence=json.loads((ROOT/'templates/release_evidence.json').read_text());self.assertTrue(workflow_bootstrap_failures({'workflow_version':'3.8.0'},evidence))
    def test_missing_render_file_blocks(self):
        (self.run/'page.png').unlink();self.assertNotEqual(run_gate(self.run)[0],0)
    def test_stale_clean_render_blocks(self):
        path=self.run/'clean-render.json';r=json.loads(path.read_text());r['source_docx_sha256']='0'*64;save(path,r);self.assertNotEqual(run_gate(self.run)[0],0)
    def test_missing_stage_blocks(self):
        path=self.run/'workflow-state.json';state=json.loads(path.read_text());state['stages'].pop();save(path,state);self.assertNotEqual(run_gate(self.run)[0],0)
    def test_human_approval_is_not_ai_review(self):
        path=self.run/'manifest.json';cfg=json.loads(path.read_text());cfg['human_approval_required']=True;save(path,cfg);self.assertNotEqual(run_gate(self.run)[0],0)

if __name__=='__main__':unittest.main()
